import pytest
from io import BytesIO
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from src.application.parsers.xlsx_manifest_reader import (
    NBSP,
    _norm_text,
    _norm_header,
    _clean,
    norm_wav_key,
    parse_xlsx_manifest,
)


@pytest.mark.unit
@pytest.mark.parametrize(
    "s,expected",
    [
        (" a  b ", "a b"),
        (f" a{NBSP}b ", "a b"),
        ("a\tb", "a b"),
        ("   ", ""),
    ],
)
def test_norm_text_normalizes_whitespace_and_nbsp(s, expected):
    assert _norm_text(s) == expected


@pytest.mark.unit
@pytest.mark.parametrize(
    "h,expected",
    [
        (None, ""),
        (" Wav_Name ", "wav_name"),
        (f"Wav{NBSP}Name", "wav name"),
        (123, "123"),
    ],
)
def test_norm_header_casefolds_and_strips(h, expected):
    assert _norm_header(h) == expected


@pytest.mark.unit
@pytest.mark.parametrize(
    "v,expected",
    [
        (None, None),
        ("", None),
        ("   ", None),
        (f"{NBSP}", None),
        ("NaN", None),
        ("none", None),
        ("NULL", None),
        (" x ", "x"),
        (10, 10),
    ],
)
def test_clean_converts_nullish_strings_to_none(v, expected):
    assert _clean(v) == expected


@pytest.mark.unit
@pytest.mark.parametrize(
    "v,expected",
    [
        (None, None),
        ("", None),
        ("  a.wav  ", "a.wav"),
        ("C:\\x\\a.wav", "a.wav"),
        ("x/y/a.wav", "a.wav"),
        (f" x{NBSP}y{NBSP}a.wav ", "x y a.wav"),
    ],
)
def test_norm_wav_key_extracts_filename_and_normalizes_spaces(v, expected):
    assert norm_wav_key(v) == expected


def _wb_bytes(*rows, add_table=False, table_ref="A1:C3"):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    if add_table:
        t = Table(displayName="T1", ref=table_ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(t)
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@pytest.mark.unit
def test_parse_xlsx_manifest_returns_empty_dict_for_empty_workbook():
    wb = Workbook()
    out = BytesIO()
    wb.save(out)
    out.seek(0)

    assert parse_xlsx_manifest(out) == {}


@pytest.mark.unit
def test_parse_xlsx_manifest_parses_header_and_data_rows_with_canonical_keys():
    xlsx = _wb_bytes(
        ("Wav_Name", "Csv_Name", "Conversation_ID", "Duration", "Xlsx_Name", "Consecutivo"),
        ("a.wav", "a.csv", "cv1", "00:00:10", "m.xlsx", "1"),
    )

    m = parse_xlsx_manifest(xlsx)

    assert "a.wav" in m
    assert m["a.wav"]["Wav_Name"] == "a.wav"
    assert m["a.wav"]["Csv_Name"] == "a.csv"
    assert m["a.wav"]["Conversation_ID"] == "cv1"
    assert m["a.wav"]["Duration"] == "00:00:10"
    assert m["a.wav"]["Xlsx_Name"] == "m.xlsx"
    assert m["a.wav"]["Consecutivo"] == "1"


@pytest.mark.unit
def test_parse_xlsx_manifest_applies_header_aliases_and_normalizes_key_from_windows_path():
    xlsx = _wb_bytes(
        ("filename", "csv name", "conversation id", "duración"),
        ("C:\\data\\A.wav", "a.csv", "cv1", "00:00:10"),
    )

    m = parse_xlsx_manifest(xlsx)

    assert "A.wav" in m
    assert m["A.wav"]["Csv_Name"] == "a.csv"
    assert m["A.wav"]["Conversation_ID"] == "cv1"


@pytest.mark.unit
def test_parse_xlsx_manifest_merges_duplicate_wav_keys_preferring_non_none_values():
    xlsx = _wb_bytes(
        ("wav_name", "csv_name", "conversation_id"),
        ("a.wav", None, "cv1"),
        ("a.wav", "a.csv", None),
    )

    m = parse_xlsx_manifest(xlsx)

    assert m["a.wav"]["Conversation_ID"] == "cv1"
    assert m["a.wav"]["Csv_Name"] == "a.csv"


@pytest.mark.unit
def test_parse_xlsx_manifest_ignores_rows_with_no_wav_name():
    xlsx = _wb_bytes(
        ("wav_name", "csv_name"),
        (None, "a.csv"),
        ("", "b.csv"),
        ("   ", "c.csv"),
    )

    assert parse_xlsx_manifest(xlsx) == {}


@pytest.mark.unit
def test_parse_xlsx_manifest_outputs_only_non_none_fields_in_result():
    xlsx = _wb_bytes(
        ("wav_name", "csv_name", "conversation_id"),
        ("a.wav", "a.csv", None),
    )

    m = parse_xlsx_manifest(xlsx)

    assert m["a.wav"]["Wav_Name"] == "a.wav"
    assert m["a.wav"]["Csv_Name"] == "a.csv"
    assert "Conversation_ID" not in m["a.wav"]
